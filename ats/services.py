import io
import re
from collections import Counter

from docx import Document
from pypdf import PdfReader


STOP_WORDS = {
    "about", "after", "also", "and", "are", "but", "for", "from", "have",
    "into", "its", "more", "our", "that", "the", "their", "this", "will",
    "with", "you", "your", "work", "working", "years", "year", "role",
    "team", "job", "candidate", "required", "requirements", "responsibilities",
    "looking", "ability", "strong", "preferred", "knowledge", "using", "must",
}

SECTION_TITLES = {
    "summary": {
        "summary", "professional summary", "profile", "about me", "executive summary",
        "career profile", "professional profile", "objective", "career objective",
    },
    "experience": {
        "experience", "work experience", "professional experience", "employment",
        "employment history", "work history", "professional background", "internship",
        "internships", "career history",
    },
    "education": {
        "education", "academic background", "qualifications", "academic qualifications",
        "coursework", "educational details", "academic record",
    },
    "skills": {
        "skills", "technical skills", "core competencies", "key skills",
        "technical proficiencies", "tech stack", "technologies", "skills & tools",
    },
    "projects": {
        "projects", "selected projects", "academic projects", "personal projects",
        "technical projects", "key projects",
    },
    "certifications": {
        "certifications", "licenses", "licenses and certifications",
        "certifications & courses", "professional certifications",
    },
    "achievements": {
        "achievements", "awards", "honors", "awards & achievements",
        "accomplishments", "key achievements",
    },
    "leadership": {
        "leadership", "extracurricular", "extracurricular activities",
        "volunteer", "volunteering", "co-curricular activities",
    },
    "languages": {
        "languages", "language proficiency", "languages & interests",
        "interests", "hobbies",
    },
}

TWELVE_STANDARD_SECTIONS = [
    {
        "id": "contact_info",
        "title": "Contact Information",
        "icon": "fa-address-card",
        "required": True,
        "aliases": {"contact", "contact information", "contact info", "personal details", "personal info"},
    },
    {
        "id": "professional_summary",
        "title": "Professional Summary",
        "icon": "fa-user-tie",
        "required": True,
        "aliases": {"summary", "professional summary", "profile", "about me", "executive summary", "career profile", "professional profile"},
    },
    {
        "id": "objective",
        "title": "Career Objective",
        "icon": "fa-bullseye",
        "required": False,
        "aliases": {"objective", "career objective", "job objective", "professional objective"},
    },
    {
        "id": "skills",
        "title": "Technical Skills",
        "icon": "fa-code",
        "required": True,
        "aliases": {"skills", "technical skills", "core competencies", "key skills", "technical proficiencies", "tech stack", "technologies", "skills & tools"},
    },
    {
        "id": "experience",
        "title": "Work Experience",
        "icon": "fa-briefcase",
        "required": True,
        "aliases": {"experience", "work experience", "professional experience", "employment", "employment history", "work history", "professional background", "career history"},
    },
    {
        "id": "internships",
        "title": "Internships",
        "icon": "fa-graduation-cap",
        "required": False,
        "aliases": {"internship", "internships", "internship experience", "industrial training"},
    },
    {
        "id": "projects",
        "title": "Projects",
        "icon": "fa-diagram-project",
        "required": True,
        "aliases": {"projects", "selected projects", "academic projects", "personal projects", "technical projects", "key projects"},
    },
    {
        "id": "education",
        "title": "Education",
        "icon": "fa-school",
        "required": True,
        "aliases": {"education", "academic background", "qualifications", "academic qualifications", "coursework", "educational details", "academic record"},
    },
    {
        "id": "certifications",
        "title": "Certifications",
        "icon": "fa-certificate",
        "required": False,
        "aliases": {"certifications", "licenses", "licenses and certifications", "certifications & courses", "professional certifications"},
    },
    {
        "id": "achievements",
        "title": "Achievements & Awards",
        "icon": "fa-trophy",
        "required": False,
        "aliases": {"achievements", "awards", "honors", "awards & achievements", "accomplishments", "key achievements"},
    },
    {
        "id": "leadership",
        "title": "Leadership & Extracurricular",
        "icon": "fa-users-gear",
        "required": False,
        "aliases": {"leadership", "extracurricular", "extracurricular activities", "volunteer", "volunteering", "co-curricular activities"},
    },
    {
        "id": "languages",
        "title": "Languages & Interests",
        "icon": "fa-language",
        "required": False,
        "aliases": {"languages", "language proficiency", "languages & interests", "interests", "hobbies"},
    },
]

EDUCATION_TERMS = {
    "bachelor", "bachelors", "bachelor's", "master", "masters", "master's",
    "degree", "diploma", "phd", "doctorate", "graduate", "postgraduate",
    "b.tech", "b.e", "b.sc", "bca", "m.tech", "m.sc", "mca",
}

SKILL_TAXONOMY = {
    "Programming Languages": [
        "python", "javascript", "typescript", "java", "c++", "c#", "c", "golang",
        "go", "rust", "ruby", "php", "swift", "kotlin", "scala", "r", "dart", "bash", "shell",
    ],
    "Frameworks & Libraries": [
        "django", "flask", "fastapi", "react", "react.js", "reactjs", "angular",
        "vue", "vue.js", "vuejs", "next.js", "nextjs", "node.js", "nodejs",
        "express", "express.js", "spring boot", "spring", "asp.net", ".net",
        "laravel", "ruby on rails", "rails", "pytorch", "tensorflow", "pandas",
        "numpy", "scikit-learn", "keras",
    ],
    "Databases": [
        "postgresql", "postgres", "mysql", "sqlite", "sqlite3", "mongodb",
        "redis", "oracle", "ms sql server", "sql server", "cassandra",
        "dynamodb", "elasticsearch", "firebase", "mariadb", "sql",
    ],
    "Frontend": [
        "html", "html5", "css", "css3", "tailwind", "tailwind css", "bootstrap",
        "sass", "scss", "webpack", "vite", "redux", "jquery", "responsive design",
    ],
    "Backend & APIs": [
        "rest api", "rest apis", "restful api", "restful apis", "graphql",
        "microservices", "celery", "grpc", "websockets", "rabbitmq", "kafka",
        "oauth", "jwt",
    ],
    "Cloud & DevOps": [
        "docker", "kubernetes", "aws", "amazon web services", "azure", "gcp",
        "google cloud platform", "google cloud", "ci/cd", "github actions",
        "gitlab ci", "jenkins", "terraform", "ansible", "linux", "nginx",
        "apache", "git", "github", "gitlab",
    ],
    "Tools & Methodologies": [
        "postman", "jira", "agile", "scrum", "unit testing", "pytest",
        "jest", "tdd", "system design", "clean code",
    ],
}

CANONICAL_SKILL_NAMES = {
    "python": "Python",
    "javascript": "JavaScript",
    "typescript": "TypeScript",
    "java": "Java",
    "c++": "C++",
    "c#": "C#",
    "c": "C",
    "go": "Go",
    "golang": "Go",
    "rust": "Rust",
    "ruby": "Ruby",
    "php": "PHP",
    "swift": "Swift",
    "kotlin": "Kotlin",
    "scala": "Scala",
    "r": "R",
    "dart": "Dart",
    "bash": "Bash",
    "shell": "Shell Scripting",
    "django": "Django",
    "flask": "Flask",
    "fastapi": "FastAPI",
    "react": "React",
    "react.js": "React",
    "reactjs": "React",
    "angular": "Angular",
    "vue": "Vue.js",
    "vue.js": "Vue.js",
    "vuejs": "Vue.js",
    "next.js": "Next.js",
    "nextjs": "Next.js",
    "node.js": "Node.js",
    "nodejs": "Node.js",
    "express": "Express.js",
    "express.js": "Express.js",
    "spring": "Spring",
    "spring boot": "Spring Boot",
    "asp.net": "ASP.NET",
    ".net": ".NET",
    "laravel": "Laravel",
    "ruby on rails": "Ruby on Rails",
    "rails": "Ruby on Rails",
    "pytorch": "PyTorch",
    "tensorflow": "TensorFlow",
    "pandas": "Pandas",
    "numpy": "NumPy",
    "scikit-learn": "Scikit-Learn",
    "keras": "Keras",
    "postgresql": "PostgreSQL",
    "postgres": "PostgreSQL",
    "mysql": "MySQL",
    "sqlite": "SQLite",
    "sqlite3": "SQLite",
    "mongodb": "MongoDB",
    "redis": "Redis",
    "oracle": "Oracle",
    "ms sql server": "MS SQL Server",
    "sql server": "SQL Server",
    "cassandra": "Cassandra",
    "dynamodb": "DynamoDB",
    "elasticsearch": "Elasticsearch",
    "firebase": "Firebase",
    "mariadb": "MariaDB",
    "sql": "SQL",
    "html": "HTML5",
    "html5": "HTML5",
    "css": "CSS3",
    "css3": "CSS3",
    "tailwind": "Tailwind CSS",
    "tailwind css": "Tailwind CSS",
    "bootstrap": "Bootstrap",
    "sass": "Sass",
    "scss": "SCSS",
    "webpack": "Webpack",
    "vite": "Vite",
    "redux": "Redux",
    "jquery": "jQuery",
    "rest api": "REST APIs",
    "rest apis": "REST APIs",
    "restful api": "RESTful APIs",
    "restful apis": "RESTful APIs",
    "graphql": "GraphQL",
    "microservices": "Microservices",
    "celery": "Celery",
    "grpc": "gRPC",
    "websockets": "WebSockets",
    "rabbitmq": "RabbitMQ",
    "kafka": "Kafka",
    "oauth": "OAuth 2.0",
    "jwt": "JWT",
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "aws": "AWS",
    "amazon web services": "AWS",
    "azure": "Microsoft Azure",
    "gcp": "Google Cloud (GCP)",
    "google cloud platform": "Google Cloud (GCP)",
    "google cloud": "Google Cloud (GCP)",
    "ci/cd": "CI/CD",
    "github actions": "GitHub Actions",
    "gitlab ci": "GitLab CI",
    "jenkins": "Jenkins",
    "terraform": "Terraform",
    "ansible": "Ansible",
    "linux": "Linux",
    "nginx": "Nginx",
    "apache": "Apache",
    "git": "Git",
    "github": "GitHub",
    "gitlab": "GitLab",
    "postman": "Postman",
    "jira": "Jira",
    "agile": "Agile Methodologies",
    "scrum": "Scrum",
    "unit testing": "Unit Testing",
    "pytest": "PyTest",
    "jest": "Jest",
    "tdd": "TDD",
    "system design": "System Design",
    "clean code": "Clean Code",
    "responsive design": "Responsive Design",
}

TECH_CASING_CORRECTIONS = {
    "python": "Python",
    "django": "Django",
    "javascript": "JavaScript",
    "typescript": "TypeScript",
    "react": "React",
    "reactjs": "React",
    "nodejs": "Node.js",
    "html": "HTML5",
    "css": "CSS3",
    "sql": "SQL",
    "postgresql": "PostgreSQL",
    "postgres": "PostgreSQL",
    "mysql": "MySQL",
    "mongodb": "MongoDB",
    "redis": "Redis",
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "aws": "AWS",
    "github": "GitHub",
    "gitlab": "GitLab",
    "api": "API",
    "apis": "APIs",
    "graphql": "GraphQL",
    "linux": "Linux",
    "postman": "Postman",
}

BUZZWORD_PATTERNS = [
    r"\bhardworking\b",
    r"\bhard-working\b",
    r"\bteam player\b",
    r"\bquick learner\b",
    r"\bfast learner\b",
    r"\bdetail-oriented\b",
    r"\bresults-driven\b",
    r"\bgo-getter\b",
    r"\bself-motivated\b",
    r"\bdynamic professional\b",
    r"\bthink outside the box\b",
]


class ResumeTextError(ValueError):
    """Raised when an uploaded resume cannot be read as searchable text."""


def extract_resume_text(resume):
    extension = resume.file_extension()
    try:
        with resume.resume_file.open("rb") as resume_file:
            file_bytes = resume_file.read()
    except Exception as error:
        raise ResumeTextError("We could not read this resume. Upload a valid PDF or DOCX file.") from error
    return extract_resume_file_text(file_bytes, extension)


def extract_resume_upload_text(uploaded_file):
    extension = uploaded_file.name.rsplit(".", 1)[-1].lower() if "." in uploaded_file.name else ""
    try:
        file_bytes = uploaded_file.read()
        uploaded_file.seek(0)
    except Exception as error:
        raise ResumeTextError("We could not read this resume. Upload a valid PDF or DOCX file.") from error
    return extract_resume_file_text(file_bytes, extension)


def extract_resume_file_text(file_bytes, extension):
    if extension not in {"pdf", "docx"}:
        raise ResumeTextError("ATS analysis supports searchable PDF and DOCX files. Legacy DOC files are not supported.")

    try:
        if extension == "pdf":
            reader = PdfReader(io.BytesIO(file_bytes), strict=False)
            if reader.is_encrypted:
                try:
                    if not reader.decrypt(""):
                        raise ResumeTextError("This PDF is password protected. Upload an unlocked copy to analyze it.")
                except NotImplementedError as error:
                    raise ResumeTextError("This PDF encryption is not supported. Upload an unlocked copy to analyze it.") from error
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(io.BytesIO(file_bytes))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            text = "\n".join(paragraphs)
    except ResumeTextError:
        raise
    except Exception as error:
        raise ResumeTextError("We could not read this resume. Upload a valid PDF or DOCX file.") from error

    text = "\n".join(
        " ".join(line.split())
        for line in text.splitlines()
        if line.strip()
    ).strip()
    if len(text) < 30:
        raise ResumeTextError("No readable text was found. Scanned PDFs need OCR before they can be analyzed.")
    return text


def _contains_phrase(text, phrase):
    escaped = re.escape(phrase.strip())
    if not escaped:
        return False
    return re.search(r"(?<!\w)" + escaped + r"(?!\w)", text, re.IGNORECASE) is not None


def _section_titles_in(resume_text):
    return {
        category
        for line in resume_text.splitlines()
        for category, titles in SECTION_TITLES.items()
        if line.strip().casefold().rstrip(":") in titles
    }


def _job_skills(job):
    if not job or not getattr(job, "skills", None):
        return []
    values = re.split(r"[,;\n]+", job.skills)
    unique = {}
    for value in values:
        cleaned = value.strip()
        if cleaned:
            unique.setdefault(cleaned.casefold(), cleaned)
    return list(unique.values())


def _job_keywords(job):
    if not job:
        return []
    job_text = " ".join(
        [
            getattr(job, "title", "") or "",
            getattr(job, "skills", "") or "",
            getattr(job, "description", "") or "",
            getattr(job, "requirements", "") or "",
            getattr(job, "responsibilities", "") or "",
        ]
    ).lower()
    words = re.findall(r"[a-z][a-z0-9+#.-]{2,}", job_text)
    counts = Counter(word.strip(".-") for word in words)
    return [
        word
        for word, _count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))
        if len(word.strip(".-")) > 2 and word.strip(".-") not in STOP_WORDS
    ][:30]


def _extract_skills_from_text(text):
    normalized = text.casefold()
    found = []
    for category, skills in SKILL_TAXONOMY.items():
        for skill in skills:
            if _contains_phrase(normalized, skill):
                canonical = CANONICAL_SKILL_NAMES.get(skill, skill.title())
                if canonical not in found:
                    found.append(canonical)
    return found


def _categorize_skills(found_skills):
    categorized = {
        "Programming Languages": [],
        "Frameworks & Libraries": [],
        "Databases": [],
        "Frontend": [],
        "Backend & APIs": [],
        "Cloud & DevOps": [],
        "Tools & Methodologies": [],
        "Other Technical Skills": [],
    }

    found_lower = {s.casefold(): s for s in found_skills}
    assigned = set()

    for category, skills in SKILL_TAXONOMY.items():
        for skill in skills:
            if skill in found_lower:
                canonical = CANONICAL_SKILL_NAMES.get(skill, found_lower[skill])
                if canonical not in categorized[category]:
                    categorized[category].append(canonical)
                assigned.add(skill)

    for skill_lower, original_name in found_lower.items():
        if skill_lower not in assigned:
            if original_name not in categorized["Other Technical Skills"]:
                categorized["Other Technical Skills"].append(original_name)

    return {k: v for k, v in categorized.items() if v}


def _segment_resume_into_sections(resume_text):
    lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
    segments = {}
    current_section = "contact_info"
    segments[current_section] = []

    # Map alias to section id
    alias_map = {}
    for sec in TWELVE_STANDARD_SECTIONS:
        for alias in sec["aliases"]:
            alias_map[alias] = sec["id"]

    for line in lines:
        line_clean = line.casefold().rstrip(":")
        if line_clean in alias_map:
            current_section = alias_map[line_clean]
            if current_section not in segments:
                segments[current_section] = []
        else:
            segments[current_section].append(line)

    return segments


def _analyze_line_improvements(segments, candidate_skills, target_keywords):
    improvements = []
    section_name_map = {sec["id"]: sec["title"] for sec in TWELVE_STANDARD_SECTIONS}

    for section_id, lines in segments.items():
        section_display = section_name_map.get(section_id, section_id.replace("_", " ").title())

        for line in lines:
            line_str = line.strip()
            if not line_str or len(line_str) < 5:
                continue

            # Check 1: First-person pronouns
            pronoun_match = re.search(r"\b(I|me|my|mine|we|our|ours)\b", line_str, re.IGNORECASE)
            if pronoun_match and section_id in {"experience", "projects", "internships", "professional_summary"}:
                # Grounded rewrite: strip leading pronoun or "I worked on / I developed"
                improved = line_str
                improved = re.sub(r"^(?:I\s+was\s+responsible\s+for|I\s+helped\s+to|I\s+worked\s+on)\s+", "Developed ", improved, flags=re.IGNORECASE)
                improved = re.sub(r"^(?:My\s+responsibilities\s+included|My\s+role\s+was\s+to)\s+", "Engineered ", improved, flags=re.IGNORECASE)
                improved = re.sub(r"^I\s+(?:am\s+a\s+|have\s+been\s+a\s+)", "", improved, flags=re.IGNORECASE)
                improved = re.sub(r"\b(?:I|we)\s+([a-zA-Z]+ed)\b", r"\1", improved, flags=re.IGNORECASE)
                improved = re.sub(r"\bmy\s+", "", improved, flags=re.IGNORECASE)
                improved = improved[0].upper() + improved[1:] if improved else line_str

                # Ensure it starts with action verb if in experience
                if section_id in {"experience", "projects"} and not re.match(r"^[A-Z][a-z]+ed\b", improved):
                    if improved.lower().startswith("developed"):
                        pass
                    else:
                        improved = "Delivered " + improved[0].lower() + improved[1:]

                matched_missing = [kw for kw in target_keywords[:4] if kw.lower() not in line_str.lower()]

                improvements.append({
                    "section": section_display,
                    "current_line": line_str,
                    "problem": "Uses first-person pronouns ('" + pronoun_match.group(0) + "'). Professional ATS formatting expects strong, third-person action verbs.",
                    "change_to": improved,
                    "why_change_it": "Demonstrates executive ownership and aligns with standard ATS parsing conventions that prioritize direct action verbs.",
                    "missing_keywords": matched_missing[:3],
                })
                continue

            # Check 2: Passive responsibility openers
            passive_match = re.search(r"^[-*•]?\s*(?:responsible for|was responsible for|helped in|assisted with|worked on|tasked with)\b", line_str, re.IGNORECASE)
            if passive_match and section_id in {"experience", "projects", "internships"}:
                improved = re.sub(r"^[-*•]?\s*(?:responsible for|was responsible for|helped in|assisted with|worked on|tasked with)\s*", "", line_str, flags=re.IGNORECASE)
                verb = "Engineered" if "api" in improved.lower() or "backend" in improved.lower() or "system" in improved.lower() else "Implemented"
                improved = f"{verb} {improved[0].lower()}{improved[1:]}"

                matched_missing = [kw for kw in target_keywords[:4] if kw.lower() not in line_str.lower()]

                improvements.append({
                    "section": section_display,
                    "current_line": line_str,
                    "problem": "Starts with passive, task-oriented phrasing ('" + passive_match.group(0).strip() + "') instead of an impactful technical action verb.",
                    "change_to": improved,
                    "why_change_it": "Action-oriented verbs highlight tangible contribution and score significantly higher on ATS semantic parsers.",
                    "missing_keywords": matched_missing[:3],
                })
                continue

            # Check 3: Generic Buzzwords
            buzz_found = []
            for bz in BUZZWORD_PATTERNS:
                m = re.search(bz, line_str, re.IGNORECASE)
                if m:
                    buzz_found.append(m.group(0))

            if buzz_found:
                improved = line_str
                for bz in buzz_found:
                    improved = re.sub(re.escape(bz), "", improved, flags=re.IGNORECASE)
                improved = re.sub(r"\s+", " ", improved).strip(" ,.-")
                if not improved:
                    improved = f"Proven track record demonstrating core technical expertise in {', '.join(candidate_skills[:3])}."

                improvements.append({
                    "section": section_display,
                    "current_line": line_str,
                    "problem": f"Contains generic subjective buzzwords ({', '.join(buzz_found)}) that lack verifiable technical proof.",
                    "change_to": improved,
                    "why_change_it": "ATS screening algorithms and technical recruiters disregard self-proclaimed adjectives in favor of hard technical competencies.",
                    "missing_keywords": [kw for kw in target_keywords[:3] if kw.lower() not in line_str.lower()],
                })
                continue

            # Check 4: Technology Casing Inaccuracies
            casing_issues = []
            fixed_line = line_str
            for lower_tech, proper_tech in TECH_CASING_CORRECTIONS.items():
                pattern = r"(?<!\w)" + re.escape(lower_tech) + r"(?!\w)"
                matches = list(re.finditer(pattern, line_str))
                for match in matches:
                    actual = match.group(0)
                    if actual != proper_tech and actual.lower() == lower_tech:
                        casing_issues.append((actual, proper_tech))
                        fixed_line = fixed_line[:match.start()] + proper_tech + fixed_line[match.end():]

            if casing_issues:
                issue_desc = ", ".join([f"'{act}' -> '{prop}'" for act, prop in casing_issues[:3]])
                improvements.append({
                    "section": section_display,
                    "current_line": line_str,
                    "problem": f"Incorrect capitalization of technical term ({issue_desc}).",
                    "change_to": fixed_line,
                    "why_change_it": "Exact technology capitalization reflects attention to detail and ensures reliable ATS keyword entity recognition.",
                    "missing_keywords": [prop for _, prop in casing_issues[:3]],
                })
                continue

            # Check 5: Career objective statement in objective section
            if section_id == "objective" and len(line_str) > 20:
                skills_preview = ", ".join(candidate_skills[:3]) if candidate_skills else "full-stack development"
                improvements.append({
                    "section": "Career Objective",
                    "current_line": line_str,
                    "problem": "Career objectives focus on personal desires rather than employer value. Modern ATS guidelines consider objectives outdated.",
                    "change_to": f"Technical Professional with established proficiency in {skills_preview}. Focused on building scalable applications and driving system reliability.",
                    "why_change_it": "Replacing an objective with an impact-focused Professional Summary elevates your profile and increases keyword relevance.",
                    "missing_keywords": [kw for kw in target_keywords[:3] if kw.lower() not in line_str.lower()],
                })
                continue

            # Check 6: Personal attributes / references clause
            if re.search(r"\b(references available upon request|marital status|date of birth|father's name)\b", line_str, re.IGNORECASE):
                improvements.append({
                    "section": section_display,
                    "current_line": line_str,
                    "problem": "Contains outdated personal attributes or reference statements that waste resume space.",
                    "change_to": "[Remove this line to conserve space for technical achievements and project details]",
                    "why_change_it": "Modern ATS resumes strictly exclude personal demographics and reference clauses to prevent bias and keep content concise.",
                    "missing_keywords": [],
                })
                continue

    return improvements


def _analyze_twelve_sections(segments, resume_text):
    results = {}
    normalized_full = resume_text.casefold()

    for sec in TWELVE_STANDARD_SECTIONS:
        sec_id = sec["id"]
        sec_title = sec["title"]
        lines = segments.get(sec_id, [])

        # Check if present in text
        is_detected = len(lines) > 0
        if not is_detected:
            # Check by alias matching across raw lines
            for alias in sec["aliases"]:
                if any(line.casefold().rstrip(":") == alias for line in resume_text.splitlines()):
                    is_detected = True
                    break

        if sec_id == "contact_info":
            has_email = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", resume_text) is not None
            has_phone = re.search(r"(?:\+?\d{1,4}[-.\s]?)?(?:\(?\d{2,4}\)?[-.\s]?)?\d{3,4}[-.\s]?\d{3,4}", resume_text) is not None
            has_linkedin = "linkedin.com" in normalized_full
            has_github = "github.com" in normalized_full

            issues = []
            if not has_email:
                issues.append("Missing email address")
            if not has_phone:
                issues.append("Missing phone number")
            if not has_linkedin:
                issues.append("LinkedIn profile not linked")
            if not has_github:
                issues.append("GitHub profile not linked")

            if issues:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "warning" if has_email and has_phone else "critical",
                    "feedback": f"Contact details found, but: {', '.join(issues)}.",
                    "issues": issues,
                    "line_count": len(lines),
                }
            else:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "ok",
                    "feedback": "No major issue detected. Contact details including email, phone, and professional profiles are present.",
                    "issues": [],
                    "line_count": len(lines),
                }

        elif sec_id == "objective":
            if is_detected:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "warning",
                    "feedback": "Career Objective section detected. Modern ATS best practices recommend replacing this with a targeted Professional Summary.",
                    "issues": ["Replace Career Objective with a Professional Summary"],
                    "line_count": len(lines),
                }
            else:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "ok",
                    "feedback": "No major issue detected. (Modern resumes prefer a Professional Summary over an Objective).",
                    "issues": [],
                    "line_count": 0,
                }

        elif not is_detected:
            if sec["required"]:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "missing",
                    "feedback": f"Section '{sec_title}' is missing. ATS algorithms heavily weight this core section.",
                    "issues": [f"Add a clearly labeled '{sec_title}' heading and corresponding details"],
                    "line_count": 0,
                }
            else:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "ok",
                    "feedback": "No major issue detected. (Optional section)",
                    "issues": [],
                    "line_count": 0,
                }
        else:
            # Section is present, check for weaknesses
            has_first_person = any(re.search(r"\b(I|me|my)\b", l, re.IGNORECASE) for l in lines)
            has_passive = any(re.search(r"^(?:responsible for|worked on)\b", l.strip(), re.IGNORECASE) for l in lines)

            issues = []
            if has_first_person:
                issues.append("Contains first-person language ('I', 'my')")
            if has_passive:
                issues.append("Contains passive bullet points starting with 'responsible for'")

            if issues:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "warning",
                    "feedback": f"Review needed: {', '.join(issues)}.",
                    "issues": issues,
                    "line_count": len(lines),
                }
            else:
                results[sec_id] = {
                    "title": sec_title,
                    "icon": sec["icon"],
                    "status": "ok",
                    "feedback": "No major issue detected. This section is well-structured and ATS-compliant.",
                    "issues": [],
                    "line_count": len(lines),
                }

    return results


def _analyze_formatting_and_parsing(resume_text):
    issues = []
    normalized = resume_text.casefold()

    # 1. Contact Info check
    if not re.search(r"[\w\.-]+@[\w\.-]+\.\w+", resume_text):
        issues.append({
            "category": "Contact Details",
            "severity": "High",
            "issue": "Missing contact email address in resume body.",
            "fix": "Place your professional email address prominently in the top contact section.",
        })
    if not re.search(r"(?:\+?\d{1,4}[-.\s]?)?(?:\(?\d{2,4}\)?[-.\s]?)?\d{3,4}[-.\s]?\d{3,4}", resume_text):
        issues.append({
            "category": "Contact Details",
            "severity": "High",
            "issue": "Missing readable phone number in resume body.",
            "fix": "Add a standard telephone number with country code in the top header.",
        })
    if "linkedin.com" not in normalized:
        issues.append({
            "category": "Online Profiles",
            "severity": "Medium",
            "issue": "Missing LinkedIn profile hyperlink.",
            "fix": "Include your customized LinkedIn profile URL (e.g. linkedin.com/in/yourname).",
        })
    if "github.com" not in normalized:
        issues.append({
            "category": "Online Profiles",
            "severity": "Low",
            "issue": "Missing GitHub or online portfolio URL.",
            "fix": "Include a link to your public GitHub profile to showcase code samples.",
        })

    # 2. Date format consistency check
    dates_found = re.findall(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{4}\b|\b\d{1,2}/\d{4}\b|\b\d{4}\s*-\s*(?:\d{4}|Present|Current)\b", resume_text, re.IGNORECASE)
    has_slash_dates = any("/" in d for d in dates_found)
    has_text_dates = any(re.search(r"[a-zA-Z]", d) for d in dates_found)
    if has_slash_dates and has_text_dates:
        issues.append({
            "category": "Date Consistency",
            "severity": "Medium",
            "issue": "Mixed date formatting styles detected (e.g., mixing '01/2023' with 'Jan 2023').",
            "fix": "Standardize all work and education dates to a single consistent style (e.g., 'Jan 2023 - Present').",
        })

    # 3. Special glyphs / emojis
    emojis = re.findall(r"[\U00010000-\U0010ffff]|[\u2600-\u26ff]|[\u2700-\u27bf]|[\u25a0-\u25ff]", resume_text)
    if emojis:
        issues.append({
            "category": "Bullet Characters",
            "severity": "Medium",
            "issue": f"Non-standard symbols or emojis detected ({', '.join(set(emojis)[:4])}).",
            "fix": "Replace decorative glyphs and emojis with standard bullet characters ('•' or '-') to ensure clean text parsing.",
        })

    # 4. Length & density
    word_count = len(resume_text.split())
    if word_count < 120:
        issues.append({
            "category": "Content Density",
            "severity": "High",
            "issue": f"Resume is brief ({word_count} words). ATS engines expect sufficient technical detail.",
            "fix": "Expand upon your technical projects, core skills, and academic coursework to achieve at least 300 words.",
        })

    # 5. Outdated Personal Elements
    if re.search(r"\breferences available upon request\b", resume_text, re.IGNORECASE):
        issues.append({
            "category": "Outdated Content",
            "severity": "Low",
            "issue": "'References available upon request' phrase detected.",
            "fix": "Remove this statement. Employers will request references during the offer stage.",
        })

    # Formatting score calculation
    deductions = sum(15 if item["severity"] == "High" else 8 if item["severity"] == "Medium" else 4 for item in issues)
    formatting_score = max(40, 100 - deductions)

    return issues, formatting_score


def analyze_resume_text(resume_text, job=None, custom_job_title="", custom_job_description=""):
    normalized_resume = re.sub(r"\s+", " ", resume_text).casefold()

    # Determine skills & keywords to match
    if job:
        required_skills = _job_skills(job)
        job_keywords = _job_keywords(job)
    elif custom_job_description:
        # Extract skills & keywords from custom job description
        required_skills = _extract_skills_from_text(custom_job_description)
        words = re.findall(r"[a-z][a-z0-9+#.-]{2,}", custom_job_description.lower())
        counts = Counter(word.strip(".-") for word in words)
        job_keywords = [
            word
            for word, _count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))
            if len(word.strip(".-")) > 2 and word.strip(".-") not in STOP_WORDS
        ][:25]
    else:
        # General resume analysis: use standard engineering expectations
        required_skills = ["Python", "SQL", "Git", "REST APIs", "Problem Solving"]
        job_keywords = ["development", "software", "applications", "database", "engineering", "testing"]

    matched_skills = [skill for skill in required_skills if _contains_phrase(normalized_resume, skill)]
    missing_skills = [skill for skill in required_skills if skill not in matched_skills]

    matched_keywords = [keyword for keyword in job_keywords if _contains_phrase(normalized_resume, keyword)]
    missing_keywords = [keyword for keyword in job_keywords if keyword not in matched_keywords]

    # Calculate existing standard scores
    skills_score = round(40 * len(matched_skills) / len(required_skills)) if required_skills else 0
    keywords_score = round(20 * len(matched_keywords) / len(job_keywords)) if job_keywords else 0

    resume_sections = _section_titles_in(resume_text)
    has_experience_section = "experience" in resume_sections
    year_values = [float(value) for value in re.findall(r"\b(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b", normalized_resume)]
    stated_years = max(year_values, default=0)
    experience_targets = {"entry": 0, "mid": 2, "senior": 5}
    target_years = experience_targets.get(getattr(job, "experience_level", "mid") if job else "mid", 0)
    experience_score = 8 if has_experience_section else 0
    if target_years == 0:
        experience_score += 17 if stated_years or has_experience_section else 0
    elif stated_years:
        experience_score += round(17 * min(stated_years / target_years, 1))

    has_education_section = "education" in resume_sections
    job_req = getattr(job, "requirements", "") or ""
    job_desc = getattr(job, "description", "") or ""
    requested_education_terms = {
        term for term in EDUCATION_TERMS
        if _contains_phrase(job_req, term) or _contains_phrase(job_desc, term)
    }
    matching_education_terms = [
        term for term in requested_education_terms if _contains_phrase(normalized_resume, term)
    ]
    if requested_education_terms:
        education_score = round(10 * len(matching_education_terms) / len(requested_education_terms))
    else:
        education_score = 10 if has_education_section else 0

    section_count = len(resume_sections)
    structure_score = 5 if section_count >= 3 else 3 if section_count == 2 else 1 if section_count == 1 else 0

    overall_score = min(
        100,
        skills_score + experience_score + keywords_score + education_score + structure_score,
    )

    # Standard suggestions (backward compatible)
    suggestions = []
    if missing_skills:
        suggestions.append(f"Highlight relevant experience with: {', '.join(missing_skills[:3])}.")
    if missing_keywords:
        suggestions.append(f"Consider including relevant terms such as: {', '.join(missing_keywords[:3])}.")
    if not has_experience_section:
        suggestions.append("Add a clearly labeled experience section with role dates and duration.")
    if not has_education_section:
        suggestions.append("Add a clearly labeled education section.")
    if not suggestions:
        suggestions.append("Your resume includes evidence for the main requirements of this role.")

    # ENRICHED ATS ANALYSIS ENGINE
    # 1. Candidate skills extraction & categorization
    extracted_candidate_skills = _extract_skills_from_text(resume_text)
    skills_categorized = _categorize_skills(extracted_candidate_skills)

    # 2. Segment into standard 12 sections
    segments = _segment_resume_into_sections(resume_text)

    # 3. Line-by-line inspection engine
    line_improvements = _analyze_line_improvements(segments, extracted_candidate_skills, missing_skills or missing_keywords)

    # 4. Section-by-section health analysis
    section_analysis = _analyze_twelve_sections(segments, resume_text)

    # 5. Formatting & parsing analysis
    formatting_issues, formatting_score = _analyze_formatting_and_parsing(resume_text)

    # 6. Projects score calculation
    has_projects = "projects" in resume_sections or len(segments.get("projects", [])) > 0
    projects_score = 90 if has_projects and len(extracted_candidate_skills) > 4 else 75 if has_projects else 40

    # 7. Critical blockers & improvement checklist
    critical_issues = []
    if not has_education_section:
        critical_issues.append("Missing clearly titled Education section.")
    if not has_experience_section and not has_projects:
        critical_issues.append("Missing both Work Experience and Projects sections.")
    if missing_skills and len(matched_skills) == 0:
        critical_issues.append(f"Zero matched required skills for target role ({', '.join(missing_skills[:3])}).")
    for fi in formatting_issues:
        if fi["severity"] == "High":
            critical_issues.append(fi["issue"])

    improvement_checklist = []
    if line_improvements:
        improvement_checklist.append(f"Revise {len(line_improvements)} flagged lines to eliminate first-person pronouns, passive phrasing, or casing errors.")
    if missing_skills:
        improvement_checklist.append(f"Incorporate missing core skills: {', '.join(missing_skills[:4])} into your Skills and Experience sections.")
    if missing_keywords:
        improvement_checklist.append(f"Add high-frequency industry keywords: {', '.join(missing_keywords[:4])}.")
    for fi in formatting_issues:
        improvement_checklist.append(fi["fix"])
    if not improvement_checklist:
        improvement_checklist.append("Your resume is well-aligned. Ensure tailored customization for each application.")

    score_breakdown = {
        "skills": {"score": skills_score, "max": 40, "label": "Skills Matching"},
        "experience": {"score": experience_score, "max": 25, "label": "Experience Alignment"},
        "keywords": {"score": keywords_score, "max": 20, "label": "Keyword Density"},
        "education": {"score": education_score, "max": 10, "label": "Education Relevance"},
        "structure": {"score": structure_score, "max": 5, "label": "Document Structure"},
        "formatting": {"score": formatting_score, "max": 100, "label": "ATS Formatting Health"},
        "projects": {"score": projects_score, "max": 100, "label": "Projects Depth"},
    }

    return {
        "overall_score": overall_score,
        "skills_score": skills_score,
        "experience_score": experience_score,
        "keywords_score": keywords_score,
        "education_score": education_score,
        "structure_score": structure_score,
        "formatting_score": formatting_score,
        "projects_score": projects_score,
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "matched_keywords": matched_keywords,
        "missing_keywords": missing_keywords,
        "suggestions": suggestions,
        "line_improvements": line_improvements,
        "section_analysis": section_analysis,
        "formatting_issues": formatting_issues,
        "skills_categorized": skills_categorized,
        "critical_issues": critical_issues,
        "score_breakdown": score_breakdown,
        "improvement_checklist": improvement_checklist,
    }
