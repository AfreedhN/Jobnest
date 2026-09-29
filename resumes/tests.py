import shutil
import tempfile
from io import BytesIO

from docx import Document
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from companies.models import Company
from jobs.models import Job
from .models import Resume


class ResumeModelTest(TestCase):

    @classmethod
    def setUpClass(cls):
        cls.temp_media = tempfile.mkdtemp()
        cls._media_override = override_settings(MEDIA_ROOT=cls.temp_media)
        cls._media_override.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls._media_override.disable()
        shutil.rmtree(cls.temp_media, ignore_errors=True)

    def setUp(self):

        self.user = User.objects.create_user(
            username="testuser",
            password="testpassword123"
        )

        self.file = SimpleUploadedFile(
            "resume.pdf",
            b"Dummy PDF content",
            content_type="application/pdf"
        )

        self.resume = Resume.objects.create(
            user=self.user,
            title="Python Developer Resume",
            resume_file=self.file,
            skills="Python, Django, HTML, CSS",
            ats_score=40,
            is_primary=True
        )

    def test_resume_created(self):

        self.assertEqual(
            self.resume.title,
            "Python Developer Resume"
        )

    def test_resume_string(self):

        self.assertEqual(
            str(self.resume),
            "Python Developer Resume - testuser"
        )

    def test_file_extension(self):

        self.assertEqual(
            self.resume.file_extension(),
            "pdf"
        )


class ResumeViewTest(TestCase):

    @classmethod
    def setUpClass(cls):
        cls.temp_media = tempfile.mkdtemp()
        cls._media_override = override_settings(MEDIA_ROOT=cls.temp_media)
        cls._media_override.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls._media_override.disable()
        shutil.rmtree(cls.temp_media, ignore_errors=True)

    def setUp(self):

        self.user = User.objects.create_user(
            username="testuser",
            password="testpassword123"
        )

        self.file = SimpleUploadedFile(
            "resume.pdf",
            b"Dummy PDF content",
            content_type="application/pdf"
        )

        self.resume = Resume.objects.create(
            user=self.user,
            title="My Resume",
            resume_file=self.file,
            skills="Python, Django",
            ats_score=20
        )

    def test_resume_list_requires_login(self):

        response = self.client.get(
            reverse("resumes:resume_list")
        )

        self.assertEqual(
            response.status_code,
            302
        )

    def test_resume_list_logged_in(self):

        self.client.login(
            username="testuser",
            password="testpassword123"
        )

        response = self.client.get(
            reverse("resumes:resume_list")
        )

        self.assertEqual(
            response.status_code,
            200
        )

        self.assertContains(
            response,
            "My Resume"
        )

    def test_resume_detail(self):

        self.client.login(
            username="testuser",
            password="testpassword123"
        )

        response = self.client.get(
            reverse(
                "resumes:resume_detail",
                kwargs={
                    "pk": self.resume.pk
                }
            )
        )

        self.assertEqual(
            response.status_code,
            200
        )

    def test_resume_page_lists_user_resumes(self):

        self.client.login(
            username="testuser",
            password="testpassword123"
        )

        response = self.client.get(
            reverse("resume")
        )

        self.assertEqual(
            response.status_code,
            200
        )
        self.assertContains(
            response,
            "My Resume"
        )

    def test_primary_resume_action_uses_post(self):
        self.client.login(
            username="testuser",
            password="testpassword123"
        )

        response = self.client.post(
            reverse("set_primary_resume", kwargs={"pk": self.resume.pk})
        )

        self.assertRedirects(response, reverse("resumes:resume_list"))
        self.resume.refresh_from_db()
        self.assertTrue(self.resume.is_primary)

    def test_ats_analyzer_uses_real_resume_and_job(self):

        self.client.login(
            username="testuser",
            password="testpassword123"
        )

        company = Company.objects.create(
            name="Acme",
            industry="Software",
            location="Remote",
            verification_status="verified",
        )

        job = Job.objects.create(
            company=company,
            title="Python Developer",
            location="Remote",
            skills="Python, Django, SQL",
            description="Build backend systems.",
            requirements="Python and Django experience.",
            responsibilities="Develop API features.",
        )

        response = self.client.get(reverse("ats_analyzer"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Python Developer")

        document = Document()
        document.add_paragraph("Skills: Python, Django, SQL")
        document.add_paragraph("Experience: 3 years building backend systems")
        output = BytesIO()
        document.save(output)
        analyzer_resume = Resume.objects.create(
            user=self.user,
            resume_file=SimpleUploadedFile(
                "ats-resume.docx",
                output.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )

        response = self.client.post(
            reverse("ats_analyzer"),
            {"resume": analyzer_resume.pk, "job": job.pk}
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ATS Compatibility Report")

    def test_delete_resume_removes_resume(self):

        self.client.login(
            username="testuser",
            password="testpassword123"
        )

        response = self.client.post(
            reverse(
                "resumes:delete_resume",
                kwargs={"pk": self.resume.pk}
            )
        )

        self.assertRedirects(
            response,
            reverse("resumes:resume_list")
        )
        self.assertFalse(
            Resume.objects.filter(pk=self.resume.pk).exists()
        )

    def test_delete_resume_requires_login(self):

        response = self.client.get(
            reverse(
                "resumes:delete_resume",
                kwargs={
                    "pk": self.resume.pk
                }
            )
        )

        self.assertEqual(
            response.status_code,
            302
        )

    def test_resume_upload_page_renders_remove_resume_button(self):
        self.client.login(
            username="testuser",
            password="testpassword123"
        )

        response = self.client.get(reverse("resume_upload"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="resumeRemoveFileBtn"')
        self.assertContains(response, 'resume-remove-file-btn')
        self.assertContains(response, "fa-trash-can")
        self.assertContains(response, 'id="fileName"')

